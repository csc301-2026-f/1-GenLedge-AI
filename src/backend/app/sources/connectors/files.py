"""Parse uploads without database access or side effects.

The returned records are the future intake writer's input. HTTP callers should
return summary() rather than echoing all uploaded data to the browser.
"""
import csv
import hashlib
import io
import json
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import BinaryIO

MAX_FILE_BYTES = 25_000_000
MAX_RECORDS = 50_000
MAX_EXPANDED_BYTES = 100_000_000
SUPPORTED_FORMATS = {"csv", "tsv", "xlsx", "xls", "json", "ndjson"}
# Python's smaller default field limit would reject otherwise valid uploads.
csv.field_size_limit(MAX_FILE_BYTES)


class FileValidationError(ValueError):
    def __init__(self, message: str, code: str = "invalid_file", status: int = 422):
        super().__init__(message)
        self.code = code
        self.status = status


@dataclass
class ParsedCollection:
    name: str
    key: str
    records: list[dict]
    headers: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def summary(self) -> dict:
        return {"name": self.name, "key": self.key,
                "rowCount": len(self.records), "headers": self.headers,
                "warnings": self.warnings}


@dataclass
class ParsedFile:
    filename: str
    format: str
    size_bytes: int
    collections: list[ParsedCollection]

    def summary(self) -> dict:
        return {"status": "validated", "persisted": False,
                "filename": self.filename, "format": self.format,
                "sizeBytes": self.size_bytes,
                "collections": [item.summary() for item in self.collections]}


def collection_key(name: str) -> str:
    # A bounded safe component, NOT a globally unique physical database name.
    slug = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")[:40] or "collection"
    return f"{slug}_{hashlib.sha256(name.encode('utf-8')).hexdigest()[:16]}"


def _append(records: list[dict], record: dict, name: str) -> None:
    if len(records) >= MAX_RECORDS:
        raise FileValidationError(f"Collection '{name}' exceeds {MAX_RECORDS:,} records.", "too_many_rows")
    records.append(record)


def _headers(values: list) -> tuple[list[dict], list[str]]:
    original = ["" if value is None else str(value) for value in values]
    bases = [value.strip() or f"column_{index + 1}" for index, value in enumerate(original)]
    reserved, used = set(bases), set()
    result, warnings = [], []
    for index, (raw, base) in enumerate(zip(original, bases)):
        key, suffix = base, 2
        while key in used or (key != base and key in reserved):
            key = f"{base}_{suffix}"
            suffix += 1
        used.add(key)
        result.append({"column": index + 1, "original": raw, "key": key})
        if raw != key:
            warnings.append(f"Column {index + 1}: {raw!r} renamed to {key!r}.")
    return result, warnings


def _table(name: str, rows, *, pad_missing: bool = False) -> ParsedCollection:
    headers, warnings, records = [], [], []
    for number, values in enumerate(rows, 1):
        values = list(values)
        if all(value is None or value == "" for value in values):
            continue
        if not headers:
            headers, warnings = _headers(values)
            continue
        if pad_missing:
            if len(values) > len(headers) and all(v is None for v in values[len(headers):]):
                values = values[:len(headers)]
            values += [None] * max(0, len(headers) - len(values))
        if len(values) != len(headers):
            raise FileValidationError(f"Collection '{name}', row {number}: expected {len(headers)} columns, got {len(values)}.")
        _append(records, dict(zip((h["key"] for h in headers), values)), name)
    if not headers:
        raise FileValidationError(f"Collection '{name}' is empty or has no header row.")
    return ParsedCollection(name, collection_key(name), records, headers, warnings)


def _reject_constant(value: str):
    raise ValueError(f"Non-standard JSON number: {value}.")


def _unique_object(pairs: list[tuple]) -> dict:
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Duplicate JSON key: {key!r}.")
        result[key] = value
    return result


def _finite_float(value: str) -> float:
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("JSON number exceeds the supported numeric range.")
    return result


def _json(text: str):
    return json.loads(text, object_pairs_hook=_unique_object,
                      parse_constant=_reject_constant, parse_float=_finite_float)


def _documents(name: str, text: str, ndjson: bool) -> ParsedCollection:
    records = []
    if ndjson:
        for number, line in enumerate(io.StringIO(text), 1):
            if not line.strip():
                continue
            try:
                item = _json(line)
            except ValueError as error:
                raise FileValidationError(f"NDJSON line {number}: {error}") from error
            if not isinstance(item, dict):
                raise FileValidationError(f"NDJSON line {number} must be a JSON object.")
            _append(records, item, name)
    else:
        try:
            data = _json(text)
        except ValueError as error:
            raise FileValidationError(f"Invalid JSON: {error}") from error
        items = [data] if isinstance(data, dict) else data
        if not isinstance(items, list):
            raise FileValidationError("JSON must contain an object or an array of objects.")
        for number, item in enumerate(items, 1):
            if not isinstance(item, dict):
                raise FileValidationError(f"JSON record {number} must be an object.")
            _append(records, item, name)
    if not records:
        raise FileValidationError("The file contains no JSON records.")
    return ParsedCollection(name, collection_key(name), records)


def _excel_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise FileValidationError("Excel contains a non-finite number.")
    return value


def _xlsx(data: bytes) -> list[ParsedCollection]:
    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > MAX_EXPANDED_BYTES:
            raise FileValidationError("Expanded XLSX contents exceed 100 MB.", "expanded_file_too_large")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        collections = []
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()

            def rows():
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        if cell.data_type in {"f", "e"}:
                            raise FileValidationError(f"Worksheet '{sheet.title}', cell {cell.coordinate}: formulas and error cells must be replaced with values.")
                        values.append(_excel_value(cell.value))
                    yield values

            collections.append(_table(sheet.title, rows(), pad_missing=True))
        return collections
    finally:
        workbook.close()


def _xls(data: bytes) -> list[ParsedCollection]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=data, on_demand=True, logfile=io.StringIO())
    try:
        collections = []
        for index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(index)
            def rows():
                for row in sheet.get_rows():
                    values = []
                    for cell in row:
                        if cell.ctype == xlrd.XL_CELL_ERROR:
                            raise FileValidationError(f"Worksheet '{sheet.name}' contains an Excel error cell.")
                        value = cell.value
                        if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        values.append(_excel_value(value))
                    yield values
            collections.append(_table(sheet.name, rows(), pad_missing=True))
            workbook.unload_sheet(sheet.name)
        return collections
    finally:
        workbook.release_resources()


def parse_file(filename: str, stream: BinaryIO) -> ParsedFile:
    """Read at most the file limit plus one byte; return only after full validation."""
    filename = filename.replace("\\", "/").rsplit("/", 1)[-1]
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in SUPPORTED_FORMATS:
        raise FileValidationError("Supported file formats: CSV, TSV, XLSX, XLS, JSON, NDJSON.", "unsupported_format", 415)
    data = stream.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise FileValidationError("File exceeds the 25 MB limit (25,000,000 bytes).", "file_too_large", 413)
    if not data:
        raise FileValidationError("The uploaded file is empty.")
    name = Path(filename).stem
    try:
        if extension in {"xlsx", "xls"}:
            collections = _xlsx(data) if extension == "xlsx" else _xls(data)
        else:
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise FileValidationError("Text files must not contain NUL characters; use UTF-8 encoding.")
            if extension in {"csv", "tsv"}:
                reader = csv.reader(io.StringIO(text, newline=""), delimiter="\t" if extension == "tsv" else ",", strict=True)
                collections = [_table(name, reader)]
            else:
                collections = [_documents(name, text, extension == "ndjson")]
        if not collections:
            raise FileValidationError("The workbook contains no worksheets.")
        keys = [collection.key for collection in collections]
        if len(keys) != len(set(keys)):
            raise FileValidationError("Collection names must be unique within a source.")
        return ParsedFile(filename, extension, len(data), collections)
    except FileValidationError:
        raise
    except UnicodeDecodeError as error:
        raise FileValidationError("Text files must use UTF-8 encoding (UTF-8 BOM is supported).") from error
    except RecursionError as error:
        raise FileValidationError("The file's nesting is too deep to parse.") from error
    except Exception as error:
        # Binary workbook readers raise several library-specific corruption errors.
        # Keep their internal details out of the public API.
        if extension in {"xlsx", "xls"} or isinstance(error, (csv.Error, ValueError)):
            raise FileValidationError(f"Unable to parse {extension.upper()}; check that the file is valid and not encrypted.") from error
        raise


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Validate a source file without a database.")
    parser.add_argument("file", type=Path)
    args = parser.parse_args()
    try:
        with args.file.open("rb") as source:
            print(json.dumps(parse_file(args.file.name, source).summary(), indent=2))
    except (FileValidationError, OSError) as error:
        parser.exit(1, f"Validation failed: {error}\n")
