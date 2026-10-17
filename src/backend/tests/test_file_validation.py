import io
import json

import pytest
from openpyxl import Workbook
import xlwt

from app.sources.connectors import files
from app.sources.connectors.files import FileValidationError, parse_file


def parse(name, data):
    return parse_file(name, io.BytesIO(data.encode() if isinstance(data, str) else data))


def workbook_bytes(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


@pytest.mark.parametrize("extension,delimiter", [("csv", ","), ("tsv", "\t")])
def test_bom_duplicate_blank_and_reserved_headers(extension, delimiter):
    data = "\ufeff" + delimiter.join(["name", "name", "name_2", "", " name "]) + "\n"
    data += delimiter.join(["001", "two", "three", "four", "five"])
    result = parse("data." + extension, data).collections[0]
    assert result.records == [{"name": "001", "name_3": "two", "name_2": "three", "column_4": "four", "name_4": "five"}]
    assert result.headers[1]["original"] == "name"
    assert len(result.warnings) == 3


def test_csv_quoted_newlines_and_blanks():
    result = parse("DATA.CSV", 'id,note\r\n1,"hello,\r\nworld"\r\n\r\n')
    assert result.collections[0].records == [{"id": "1", "note": "hello,\r\nworld"}]
    assert result.summary()["persisted"] is False
    assert "records" not in result.summary()["collections"][0]


def test_large_csv_cell_within_file_limit():
    value = "x" * 200_000
    assert parse("a.csv", "note\n" + value).collections[0].records == [{"note": value}]


@pytest.mark.parametrize("name,data", [
    ("a.csv", "a,b\n1"), ("a.tsv", "a\n1\t2"), ("a.csv", 'a\n"unfinished'),
    ("a.csv", b"a\n\xff"), ("a.csv", b"a\x00b"), ("a.csv", ""),
    ("a.csv", "\n\n"), ("a.xlsx", b"not a workbook"), ("a.xls", b"not a workbook"),
    ("a.json", "[1]"), ("a.json", "null"), ("a.json", "[]"),
    ("a.json", '{"a":1,"a":2}'), ("a.json", '{"x":{"a":1,"a":2}}'),
    ("a.json", '{"x":NaN}'), ("a.json", '{"x":1e999}'),
    ("a.ndjson", '{}\n{"unfinished"'), ("a.ndjson", "{}\n[]"),
    ("a.ndjson", "\n"), ("a.exe", "a,b"),
])
def test_invalid_files(name, data):
    with pytest.raises(FileValidationError):
        parse(name, data)


@pytest.mark.parametrize("extension", ["json", "ndjson"])
def test_nested_documents_preserved(extension):
    record = {"_id": {"$oid": "abc"}, "customer": {"name": "Ada"}, "items": [{"qty": 2}], "empty": None}
    result = parse("orders." + extension, json.dumps(record))
    assert result.collections[0].records == [record]


def test_ndjson_error_has_line_number():
    with pytest.raises(FileValidationError, match="line 3"):
        parse("a.ndjson", '{}\n\n{"broken"')


@pytest.mark.parametrize("extension", ["csv", "tsv", "json", "ndjson", "xlsx", "xls"])
def test_record_limit_for_every_format(monkeypatch, extension):
    monkeypatch.setattr(files, "MAX_RECORDS", 2)
    def data(count):
        if extension in {"csv", "tsv"}:
            return "id\n" + "\n".join(str(i) for i in range(count))
        if extension == "json":
            return json.dumps([{"id": i} for i in range(count)])
        if extension == "ndjson":
            return "\n".join(json.dumps({"id": i}) for i in range(count))
        if extension == "xlsx":
            return workbook_bytes([("Rows", [["id"]] + [[i] for i in range(count)])])
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("Rows")
        sheet.write(0, 0, "id")
        for i in range(count):
            sheet.write(i + 1, 0, i)
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()
    assert len(parse("a." + extension, data(2)).collections[0].records) == 2
    with pytest.raises(FileValidationError, match="exceeds"):
        parse("a." + extension, data(3))


def test_real_csv_row_boundary():
    assert len(parse("a.csv", "id\n" + "1\n" * 50_000).collections[0].records) == 50_000
    with pytest.raises(FileValidationError, match="50,000"):
        parse("a.csv", "id\n" + "1\n" * 50_001)


def test_real_byte_boundary():
    # Spaces keep the document small after parsing while exercising the byte limit.
    data = b"{}" + b" " * (files.MAX_FILE_BYTES - 2)
    assert parse("a.json", data).size_bytes == files.MAX_FILE_BYTES
    with pytest.raises(FileValidationError) as error:
        parse("a.json", data + b" ")
    assert error.value.status == 413


def test_multisheet_names_are_safe_unique_and_stable():
    data = workbook_bytes([("Sales EU", [["id"], [1]]), ("Sales-EU", [["id"], [2]])])
    result = parse("a.xlsx", data)
    assert [c.name for c in result.collections] == ["Sales EU", "Sales-EU"]
    keys = [c.key for c in result.collections]
    assert len(set(keys)) == 2
    assert keys == [c.key for c in parse("a.xlsx", data).collections]


def test_later_bad_sheet_prevents_any_result():
    data = workbook_bytes([("Good", [["id"], [1]]), ("Bad", [["id"], ["=1+1"]])])
    with pytest.raises(FileValidationError, match="formulas"):
        parse("a.xlsx", data)


def test_xlsx_dates_missing_cells_and_empty_sheet():
    from datetime import datetime
    data = workbook_bytes([("Dates", [["date", "note"], [datetime(2026, 1, 2)]])])
    assert parse("a.xlsx", data).collections[0].records == [{"date": "2026-01-02T00:00:00", "note": None}]
    with pytest.raises(FileValidationError, match="empty"):
        parse("a.xlsx", workbook_bytes([("Empty", [])]))


def test_expansion_limit(monkeypatch):
    monkeypatch.setattr(files, "MAX_EXPANDED_BYTES", 1)
    with pytest.raises(FileValidationError, match="Expanded"):
        parse("a.xlsx", workbook_bytes([("Data", [["id"], [1]])]))


def test_filename_paths_are_removed():
    assert parse(r"C:\fakepath\data.csv", "id\n1").filename == "data.csv"
